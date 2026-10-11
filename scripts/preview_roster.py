"""Identity-only spreadsheet import and preliminary ranks from local series data."""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from core import DefaultNormalizer

ONLINE_CONTEST_IDS = (
    "icpc2026preliminary-1", "icpc2026preliminary-2", "ccpc2026preliminary",
)
ONLINE_LABELS = ("ICPC1", "ICPC2", "CCPC")


def parse_roster_rows(rows):
    """Read only identity columns; ignore every supplied score and rank."""
    from scripts.generate_preview_data import RegisteredTeam

    rows = iter(rows)
    required = ("学校或单位", "中文队名", "类型", "队员1", "队员2", "队员3", "队伍 ID")
    ccpc_headers = ("序号", "所属学校", "中文队名", "队员1姓名", "队员2姓名", "队员3姓名")
    header = None
    teams = []
    seen = set()
    for row_number, row in enumerate(rows, 1):
        if header is None:
            if all(key in row for key in ccpc_headers):
                return parse_ccpc_roster_rows(rows, row, row_number)
            if all(key in row for key in required):
                if any(row.count(key) != 1 for key in required):
                    raise ValueError(f"roster row {row_number}: duplicate identity header")
                header = {key: row.index(key) for key in required}
            continue
        if not any(value is not None for value in row):
            continue
        def value(key, header=header, row=row):
            index = header[key]
            return row[index] if index < len(row) else None
        def identity(key, value=value, row_number=row_number):
            result = value(key)
            if not isinstance(result, str) or not result.strip() or result.startswith("="):
                raise ValueError(f"roster row {row_number}: invalid {key}")
            return result.strip()
        school, name, kind, fid = (identity(key) for key in required[:3] + required[-1:])
        if fid in seen:
            raise ValueError(f"roster row {row_number}: duplicate team ID {fid}")
        seen.add(fid)
        if kind not in ("正式", "打星"):
            raise ValueError(f"roster row {row_number}: unknown team type {kind!r}")
        members = tuple(identity(key) for key in required[3:6] if value(key) not in (None, ""))
        if not members or len(set(members)) != len(members):
            raise ValueError(f"roster row {row_number}: missing or duplicate members")
        teams.append(RegisteredTeam(len(teams), name, school, members, kind == "打星"))
    if not teams:
        raise ValueError("roster: missing identity headers or empty roster")
    return teams


def parse_ccpc_roster_rows(rows, header_row, header_number):
    """CCPC export has no team type; treat entries as official, ignoring coaches."""
    keys = ("序号", "所属学校", "中文队名", "队员1姓名", "队员2姓名", "队员3姓名")
    if any(header_row.count(key) != 1 for key in keys):
        raise ValueError(f"roster row {header_number}: duplicate identity header")
    positions = [header_row.index(key) for key in keys]
    converted = [["学校或单位", "中文队名", "类型", "队员1", "队员2", "队员3", "队伍 ID"]]
    converted[:0] = [[None]] * (header_number - 1)
    for row_number, row in enumerate(rows, header_number + 1):
        if not any(value is not None for value in row):
            converted.append([None])
            continue
        serial, school, name, *members = [row[i] if i < len(row) else None for i in positions]
        if type(serial) is not int or serial < 1:
            raise ValueError(f"roster row {row_number}: invalid 序号")
        converted.append([school, name, "正式", *members, str(serial)])
    return parse_roster_rows(converted)


def load_roster_xlsx(path: Path):
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        return parse_roster_rows(workbook.worksheets[0].iter_rows(values_only=True))
    finally:
        workbook.close()


def parse_identity_roster(document):
    """Read the reviewed, identity-only JSON snapshot including eligibility."""
    teams = document.get("teams") if isinstance(document, dict) else None
    if not isinstance(teams, list) or not teams:
        raise ValueError("identity roster: expected nonempty teams")
    rows = [["学校或单位", "中文队名", "类型", "队员1", "队员2", "队员3", "队伍 ID"]]
    for index, team in enumerate(teams):
        context = f"identity roster team {index + 1}"
        if not isinstance(team, dict) or type(team.get("excluded")) is not bool:
            raise ValueError(f"{context}: expected boolean excluded")
        members = team.get("members")
        if not isinstance(members, list) or not 1 <= len(members) <= 3:
            raise ValueError(f"{context}: expected one to three members")
        if any(not isinstance(m, str) or not m.strip() for m in members):
            raise ValueError(f"{context}: invalid member")
        rows.append([team.get("school"), team.get("name"),
                     "打星" if team["excluded"] else "正式",
                     *members, *([None] * (3 - len(members))), str(index)])
    result = parse_roster_rows(rows)
    keys = {(t.school, t.name, tuple(sorted(t.members))) for t in result}
    if len(keys) != len(result):
        raise ValueError("identity roster: duplicate team identity")
    return result


def online_rank_index(series: dict, normalizer: DefaultNormalizer, cutoff: str):
    """Index official participation ranks, resolving duplicate identities by best rank."""
    selected = {}
    for index, contest in enumerate(series["contests"]):
        if contest["id"] in ONLINE_CONTEST_IDS:
            slot = ONLINE_CONTEST_IDS.index(contest["id"])
            if slot in selected.values():
                raise ValueError(f"duplicate online contest {contest['id']}")
            if datetime.fromisoformat(contest["startAt"]) >= datetime.fromisoformat(cutoff):
                raise ValueError(f"{contest['id']}: online result must precede preview")
            selected[index] = slot
    if len(selected) != 3:
        raise ValueError("online series: expected both ICPC preliminaries and CCPC preliminary")
    result = {}
    for person in series["competitors"]:
        key = (normalizer.school(person["school"]), normalizer.member(person["member"]))
        ranks = result.setdefault(key, [None, None, None])
        for participation in person["participations"]:
            index = participation["contestIndex"]
            if index not in selected:
                continue
            rank = participation["contestRank"]
            if type(rank) is not int or rank < 1:
                raise ValueError(f"online series: invalid rank for {key}: {rank!r}")
            slot = selected[index]
            ranks[slot] = rank if ranks[slot] is None else min(ranks[slot], rank)
    return result


def attach_online_ranks(teams: list[dict], index: dict, normalizer: DefaultNormalizer):
    for team in teams:
        for member in team["members"]:
            key = (normalizer.school(team["school"]), normalizer.member(member["name"]))
            member["onlineRanks"] = list(index.get(key, [None, None, None]))
        team["onlineRanks"] = [
            min((member["onlineRanks"][slot] for member in team["members"]
                 if member["onlineRanks"][slot] is not None), default=None)
            for slot in range(3)
        ]
