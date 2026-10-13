"""Generate summary-only regional cutoffs from identity columns and local online results."""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

from scripts.generate_preview_data import load_normalizer
from scripts.preview_roster import attach_online_ranks, online_rank_index, parse_roster_rows


def calculate_cutoffs(teams):
    official = [team for team in teams if not team.get("excluded", False)]
    ordered = sorted(official, key=lambda t: sorted(
        rank if rank is not None else float("inf") for rank in t["onlineRanks"]
    ))
    schools = {}
    for team in ordered:
        schools.setdefault(team["school"], team)
    n = len(official)
    gold = (n + 9) // 10
    silver = gold + (2 * n + 9) // 10
    bronze = silver + (3 * n + 9) // 10
    boundaries = {}
    for name, rows, position in [("cup", list(schools.values()), 3),
                                 ("gold", ordered, gold), ("silver", ordered, silver),
                                 ("bronze", ordered, bronze)]:
        boundaries[name] = rows[position - 1] if 0 < position <= len(rows) else None
    cutoffs = {key: min((r for r in team["onlineRanks"] if r is not None), default=None)
               if team else None for key, team in boundaries.items()}
    return cutoffs, boundaries


def build_summary(registered, series, normalizer, *, contest_id, title, sort_at, host_school):
    teams = [{"school": normalizer.school(t.school), "name": t.name,
              "members": [{"name": m} for m in t.members], "excluded": t.excluded}
             for t in registered]
    identities = {(t["school"], t["name"], tuple(sorted(m["name"] for m in t["members"])))
                  for t in teams}
    if len(identities) != len(teams):
        raise ValueError("roster: duplicate team identity")
    attach_online_ranks(teams, online_rank_index(series, normalizer, sort_at), normalizer)
    cutoffs, boundaries = calculate_cutoffs(teams)
    official = [t for t in teams if not t["excluded"]]
    summary = {
        "id": contest_id, "title": title, "contestStage": "regional", "sortAt": sort_at,
        "hostSchool": host_school, "cutoffs": cutoffs,
        "teamCount": len(teams), "officialTeamCount": len(official),
        "excludedTeamCount": len(teams) - len(official),
        "matchedOfficialTeamCount": sum(any(r is not None for r in t["onlineRanks"])
                                        for t in official),
    }
    return summary, {"teams": teams, "boundaries": boundaries}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    for key in ("teams", "series", "output", "audit-output"):
        parser.add_argument(f"--{key}", type=Path, required=True)
    for key in ("sheet", "contest-id", "title", "sort-at", "host-school", "source-url"):
        parser.add_argument(f"--{key}", required=True)
    parser.add_argument("--school-aliases", type=Path, default=Path("config/school-aliases.json"))
    args = parser.parse_args()
    from openpyxl import load_workbook

    workbook = load_workbook(args.teams, read_only=True, data_only=False)
    try:
        if args.sheet not in workbook.sheetnames:
            raise ValueError(f"missing sheet: {args.sheet}")
        registered = parse_roster_rows(workbook[args.sheet].iter_rows(values_only=True))
    finally:
        workbook.close()
    series = json.loads(args.series.read_text(encoding="utf-8"))
    summary, audit = build_summary(
        registered, series, load_normalizer(args.school_aliases), contest_id=args.contest_id,
        title=args.title, sort_at=args.sort_at, host_school=args.host_school,
    )
    summary["sources"] = {
        "rosterFile": args.teams.name, "rosterSheet": args.sheet,
        "rosterSha256": hashlib.sha256(args.teams.read_bytes()).hexdigest(),
        "onlineSeriesSha256": hashlib.sha256(args.series.read_bytes()).hexdigest(),
        "schoolAliasesSha256": hashlib.sha256(args.school_aliases.read_bytes()).hexdigest(),
        "contestUrl": args.source_url,
    }
    document = (json.loads(args.output.read_text(encoding="utf-8")) if args.output.exists()
                else {"schemaVersion": 1, "seriesId": series["id"], "contests": []})
    if document["seriesId"] != series["id"] or document["schemaVersion"] != 1:
        raise ValueError("output: incompatible regional cutoff document")
    document["contests"] = sorted(
        [c for c in document["contests"] if c["id"] != summary["id"]] + [summary],
        key=lambda c: c["id"],
    )
    for path, value in [(args.output, document), (args.audit_output, audit)]:
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=True))


if __name__ == "__main__":
    main()
